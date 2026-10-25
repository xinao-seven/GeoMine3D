"""建模引擎输入适配层的单元测试(不依赖数据库)。

覆盖:两种分层表列语义的解析、层序推导、建模输入交集、界面控制点、
保序投影与岩性配色。
"""
from pathlib import Path

import numpy as np
import pytest
from openpyxl import Workbook

from modeling_engine import data_io
from modeling_engine.isotonic import project_order
from modeling_engine.package_builder import layer_color
from modeling_engine.surfaces import build_interface_points


def _write_strata(path: Path, rows: list[tuple[str, str, float, float]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["钻孔名称", "地层名称", "深度", "厚度"])
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _write_locations(path: Path, rows: list[tuple[str, float, float, float]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "y", "x", "z"])
    for name, y, x, z in rows:
        sheet.append([name, y, x, z])
    workbook.save(path)


class TestParseStrata:
    def test_top_bottom_convention(self, tmp_path: Path):
        # 14层汇总表口径:深度=层顶,厚度列=层底(恒 ≥ 深度)
        path = tmp_path / "strata.xlsx"
        _write_strata(path, [
            ("J001", "土层", 0, 10),
            ("J001", "煤3-1", 10, 12),
            ("J002", "土层", 0, 5),
        ])
        strata = data_io.parse_strata(path)
        assert strata.convention == "top_bottom"
        assert strata.hole_count == 2
        segments = strata.holes["J001"].segments
        assert [(s["top_depth"], s["bottom_depth"]) for s in segments] == [(0, 10), (10, 12)]

    def test_bottom_thickness_convention(self, tmp_path: Path):
        # 旧表口径:深度=层底,厚度=层厚 → top = 深度 - 厚度
        path = tmp_path / "strata_legacy.xlsx"
        _write_strata(path, [
            ("J001", "土层", 10, 10),
            ("J001", "煤3-1", 12, 2),
        ])
        strata = data_io.parse_strata(path)
        assert strata.convention == "bottom_thickness"
        segments = strata.holes["J001"].segments
        assert [(s["top_depth"], s["bottom_depth"]) for s in segments] == [(0, 10), (10, 12)]

    def test_negative_thickness_rejected(self, tmp_path: Path):
        path = tmp_path / "strata_bad.xlsx"
        # 旧表语义(厚度<深度):厚度为负 → 层厚 = 厚度 < 0
        _write_strata(path, [
            ("J001", "土层", 10, 8),
            ("J001", "煤3-1", 12, -3),
        ])
        with pytest.raises(data_io.DataFormatError, match="层厚为负"):
            data_io.parse_strata(path)

    def test_missing_header_rejected(self, tmp_path: Path):
        path = tmp_path / "strata_missing.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["钻孔名称", "地层名称", "深度"])  # 缺 厚度
        sheet.append(["J001", "土层", 5])
        workbook.save(path)
        with pytest.raises(data_io.DataFormatError, match="厚度"):
            data_io.parse_strata(path)

    def test_match_key_normalization(self, tmp_path: Path):
        path = tmp_path / "strata.xlsx"
        _write_strata(path, [("j-001", "土层", 0, 10)])
        strata = data_io.parse_strata(path)
        assert "J001" in strata.holes  # 编号匹配键仅保留字母数字并大写


class TestDeriveLayers:
    def test_ordered_by_mean_top_depth(self, tmp_path: Path):
        path = tmp_path / "strata.xlsx"
        _write_strata(path, [
            ("J001", "煤层", 10, 12),
            ("J001", "土层", 0, 10),
            ("J002", "煤层", 5, 7),
            ("J002", "土层", 0, 5),
            ("J002", "底岩", 7, 20),
        ])
        strata = data_io.parse_strata(path)
        layers = data_io.derive_layers(strata.holes)
        # 土层均顶 0 < 煤层 7.5 < 底岩 7?底岩 J002 顶 7、煤层均顶 (10+5)/2=7.5 → 土层,底岩,煤层
        assert layers == ["土层", "底岩", "煤层"]

    def test_single_layer_rejected(self, tmp_path: Path):
        path = tmp_path / "strata.xlsx"
        _write_strata(path, [("J001", "土层", 0, 10)])
        strata = data_io.parse_strata(path)
        with pytest.raises(data_io.DataFormatError, match="至少需要 2 个"):
            data_io.derive_layers(strata.holes)


class TestBuildModelInputs:
    def test_intersection_and_skipped(self, tmp_path: Path):
        strata_path = tmp_path / "strata.xlsx"
        _write_strata(strata_path, [
            ("J001", "土层", 0, 10),
            ("J001", "煤层", 10, 12),
            ("J002", "土层", 0, 8),
            ("J002", "煤层", 8, 10),
            ("J003", "土层", 0, 6),
        ])
        location_path = tmp_path / "location.xlsx"
        _write_locations(location_path, [
            ("J001", 4300000, 37400000, 1200),
            ("J002", 4300100, 37400100, 1210),
            ("J009", 4300200, 37400200, 1220),
        ])
        strata = data_io.parse_strata(strata_path)
        locations = data_io.parse_locations(location_path)
        inputs = data_io.build_model_inputs(strata, locations)
        assert inputs.boreholes["borehole_id"].tolist() == ["J001", "J002"]
        assert len(inputs.intervals) == 4
        assert inputs.layers == ["土层", "煤层"]
        reasons = {item["code"]: item["reason"] for item in inputs.skipped}
        assert "位置表中没有该钻孔坐标" in reasons["J003"]
        assert "分层表中没有该钻孔" in reasons["J009"]

    def test_no_common_borehole_rejected(self, tmp_path: Path):
        strata_path = tmp_path / "strata.xlsx"
        _write_strata(strata_path, [
            ("J001", "土层", 0, 10),
            ("J001", "煤层", 10, 12),
        ])
        location_path = tmp_path / "location.xlsx"
        _write_locations(location_path, [("J009", 4300000, 37400000, 1200)])
        strata = data_io.parse_strata(strata_path)
        locations = data_io.parse_locations(location_path)
        with pytest.raises(data_io.DataFormatError, match="没有共同钻孔"):
            data_io.build_model_inputs(strata, locations)


class TestInterfacePoints:
    def test_top_bottom_elevation(self):
        import pandas as pd

        boreholes = pd.DataFrame([
            {"borehole_id": "J001", "x": 0.0, "y": 0.0, "z": 1200.0},
        ])
        intervals = pd.DataFrame([
            {"borehole_id": "J001", "layer_name": "土层", "top_depth": 0, "bottom_depth": 10},
            {"borehole_id": "J001", "layer_name": "煤层", "top_depth": 10, "bottom_depth": 12},
        ])
        interfaces, points, all_points = build_interface_points(boreholes, intervals, ["土层", "煤层"])
        assert interfaces == ["IF_0_top_土层", "IF_1_bottom_土层", "IF_2_bottom_煤层"]
        assert len(all_points) == 3
        assert points["IF_0_top_土层"]["z"].iloc[0] == 1200
        assert points["IF_1_bottom_土层"]["z"].iloc[0] == 1190
        assert points["IF_2_bottom_煤层"]["z"].iloc[0] == 1188


class TestProjectOrder:
    def test_enforces_min_thickness(self):
        z = np.array([100.0, 90.0, 95.0, 60.0])
        weights = np.ones(4)
        projected = project_order(z, weights, min_thickness=0.0)
        assert np.all(np.diff(projected) <= 1e-9)  # 单调不增
        # 违序位置(90 → 95)应被压平
        assert projected[1] >= projected[2] - 1e-9


class TestLayerColor:
    def test_known_names(self):
        assert layer_color("风积砂层") == "#D9C7A3"
        assert layer_color("煤3-1") == "#333333"

    def test_pattern_fallback(self):
        assert layer_color("第5旋回泥岩层") == "#8C6D5A"
        assert layer_color("第5旋回砂岩层") == "#D98C5F"
        assert layer_color("未知层") == "#B0A48F"
