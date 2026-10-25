"""输入适配层:钻孔 Excel → 建模引擎标准 DataFrame。

分层长表(表头 钻孔名称/地层名称/深度/厚度)存在两种列语义,表头完全相同:
- top_bottom(当前14层表): 深度=层顶深度, 厚度列实际存层底深度;
- bottom_thickness(旧表): 深度=层底深度, 厚度=层厚。
判别逻辑与 app/services/server_import_service.detect_strata_convention 一致:
全部行满足 厚度列 ≥ 深度列 时只有"顶底"写法成立。

位置表表头 name|钻孔名称/borehole_id + x/y/z(Gauss-Kruger: x=东, y=北, z=孔口高程)。

钻孔编号匹配键与后端一致:仅保留字母数字并转大写(J-001 与 J001 视为同一孔)。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook


class DataFormatError(ValueError):
    """输入 Excel 格式不满足建模要求。"""


def _read_rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        return [dict(zip(headers, row, strict=False)) for row in rows if any(row)]
    except StopIteration as exc:
        raise DataFormatError(f"{path.name}: 工作表为空") from exc
    finally:
        workbook.close()


def normalize_match_key(value: Any) -> str:
    return "".join(char for char in str(value or "").strip().upper() if char.isalnum())


def canonical_code(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).upper()


def _as_float(value: Any, field_label: str, context: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise DataFormatError(f"{context}: {field_label} 不是数值({value!r})") from exc
    if not np.isfinite(result):
        raise DataFormatError(f"{context}: {field_label} 不是有限数值({value!r})")
    return result


def _pick(row: dict[str, Any], names: tuple[str, ...]) -> Any:
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return None


@dataclass
class StrataHole:
    """单孔分层(已按层顶深度排序)。"""

    code: str
    segments: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class StrataData:
    convention: str
    holes: dict[str, StrataHole]  # 匹配键 → 孔

    @property
    def hole_count(self) -> int:
        return len(self.holes)


def detect_strata_convention(rows: list[dict[str, Any]]) -> str:
    pairs: list[tuple[float, float]] = []
    for row in rows:
        if not normalize_match_key(_pick(row, ("钻孔名称", "name", "borehole_id"))):
            continue
        first, second = row.get("深度"), row.get("厚度")
        if first in (None, "") or second in (None, ""):
            continue
        pairs.append((float(first), float(second)))
    if pairs and all(second >= first for first, second in pairs):
        return "top_bottom"
    return "bottom_thickness"


def parse_strata(path: Path) -> StrataData:
    rows = _read_rows(path)
    if not rows:
        raise DataFormatError(f"{path.name}: 分层表没有数据行")
    missing = [c for c in ("钻孔名称", "地层名称", "深度", "厚度") if c not in rows[0]]
    if missing:
        raise DataFormatError(
            f"{path.name}: 分层表缺少表头列 {', '.join(missing)}; "
            "需要 钻孔名称/地层名称/深度/厚度 四列"
        )
    convention = detect_strata_convention(rows)
    holes: dict[str, StrataHole] = {}
    for row in rows:
        raw_name = _pick(row, ("钻孔名称", "name", "borehole_id"))
        key = normalize_match_key(raw_name)
        if not key:
            continue
        layer_name = str(_pick(row, ("地层名称", "layer_name")) or "未命名地层").strip()
        if not layer_name:
            layer_name = "未命名地层"
        context = f"{path.name} 钻孔 {canonical_code(raw_name)} 地层 {layer_name}"
        first = _as_float(row.get("深度"), "深度", context)
        second = _as_float(row.get("厚度"), "厚度", context)
        if convention == "top_bottom":
            top_depth, bottom_depth = first, second
        else:
            bottom_depth, top_depth = first, first - second
        thickness = bottom_depth - top_depth
        if thickness < -1e-6:
            raise DataFormatError(
                f"{context}: 层厚为负({thickness:.3f}m)。请确认列语义是 "
                "深度=层顶+厚度=层底(14层汇总表), 还是 深度=层底+厚度=层厚(旧表)"
            )
        hole = holes.setdefault(key, StrataHole(code=canonical_code(raw_name)))
        hole.segments.append(
            {
                "layer_name": layer_name,
                "top_depth": top_depth,
                "bottom_depth": bottom_depth,
            }
        )
    if not holes:
        raise DataFormatError(f"{path.name}: 分层表中没有有效钻孔行")
    for hole in holes.values():
        hole.segments.sort(key=lambda item: item["top_depth"])
    return StrataData(convention=convention, holes=holes)


def parse_locations(path: Path) -> dict[str, dict[str, Any]]:
    """读取钻孔位置表 → 匹配键 → {code, x, y, z}(x=东, y=北, z=高程)。"""
    rows = _read_rows(path)
    if not rows:
        raise DataFormatError(f"{path.name}: 位置表没有数据行")
    name_keys = ("name", "钻孔名称", "borehole_id")
    if not any(key in rows[0] for key in name_keys) or not all(
        key in rows[0] for key in ("x", "y", "z")
    ):
        raise DataFormatError(
            f"{path.name}: 位置表需要 名称(name/钻孔名称) 与 x/y/z 列"
        )
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        raw_name = _pick(row, name_keys)
        key = normalize_match_key(raw_name)
        if not key:
            continue
        context = f"{path.name} 钻孔 {canonical_code(raw_name)}"
        result[key] = {
            "code": canonical_code(raw_name),
            "x": _as_float(row.get("x"), "x", context),
            "y": _as_float(row.get("y"), "y", context),
            "z": _as_float(row.get("z"), "z", context),
        }
    if not result:
        raise DataFormatError(f"{path.name}: 位置表中没有有效钻孔行")
    return result


def derive_layers(holes: dict[str, StrataHole]) -> list[str]:
    """层序推导:按各层平均层顶深度自上而下排序(与沉积顺序一致)。"""
    stats: dict[str, list[float]] = {}
    order_seen: list[str] = []
    for hole in holes.values():
        for segment in hole.segments:
            name = segment["layer_name"]
            if name not in stats:
                stats[name] = []
                order_seen.append(name)
            stats[name].append(segment["top_depth"])
    if len(stats) < 2:
        raise DataFormatError(
            f"分层表中只识别出 {len(stats)} 个层位,至少需要 2 个才能建模"
        )
    return sorted(stats, key=lambda name: (float(np.mean(stats[name])), order_seen.index(name)))


@dataclass
class ModelInputs:
    """引擎建模输入(与论文 chapter34_full_experiments.load_data 输出契约一致)。"""

    boreholes: pd.DataFrame      # borehole_id, x, y, z
    intervals: pd.DataFrame      # borehole_id, layer_name, top_depth, bottom_depth
    layers: list[str]
    display_codes: dict[str, str]  # 匹配键 → 展示编号
    skipped: list[dict[str, str]]


def build_model_inputs(
    strata: StrataData, locations: dict[str, dict[str, Any]]
) -> ModelInputs:
    """分层表 ∩ 位置表:只有同时具备分层与坐标的孔才能参与建模。"""
    layers = derive_layers(strata.holes)
    display_codes: dict[str, str] = {}
    borehole_rows: list[dict[str, float]] = []
    interval_rows: list[dict[str, Any]] = []
    skipped: list[dict[str, str]] = []

    for key, hole in strata.holes.items():
        location = locations.get(key)
        if location is None:
            skipped.append({"code": hole.code, "reason": "位置表中没有该钻孔坐标"})
            continue
        display_codes[key] = hole.code
        borehole_rows.append(
            {"borehole_id": key, "x": location["x"], "y": location["y"], "z": location["z"]}
        )
        for segment in hole.segments:
            if segment["layer_name"] not in layers:  # 理论不可达,防御性保留
                continue
            interval_rows.append(
                {
                    "borehole_id": key,
                    "layer_name": segment["layer_name"],
                    "top_depth": segment["top_depth"],
                    "bottom_depth": segment["bottom_depth"],
                }
            )

    if not borehole_rows:
        raise DataFormatError(
            "分层表与位置表没有共同钻孔:请检查两个文件的钻孔名称是否一致"
        )
    for key, location in locations.items():
        if key not in strata.holes:
            skipped.append({"code": location["code"], "reason": "分层表中没有该钻孔"})

    boreholes = pd.DataFrame(borehole_rows)
    intervals = pd.DataFrame(
        interval_rows, columns=["borehole_id", "layer_name", "top_depth", "bottom_depth"]
    )
    return ModelInputs(
        boreholes=boreholes,
        intervals=intervals,
        layers=layers,
        display_codes=display_codes,
        skipped=skipped,
    )
